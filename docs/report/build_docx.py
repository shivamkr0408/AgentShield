"""Build report.docx from report.md (headings, paragraphs, bullets, tables, images, bold).

    pip install python-docx
    python docs/report/build_docx.py

A pragmatic Markdown subset converter, enough for this report. Figures referenced with
![alt](figures/x.png) are embedded.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
MD = HERE / "report.md"
OUT = HERE / "report.docx"

_BOLD = re.compile(r"\*\*(.+?)\*\*")
_IMG = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")
_LINK = re.compile(r"\[(.+?)\]\((.+?)\)")


def add_runs(paragraph, text: str) -> None:
    text = _LINK.sub(r"\1", text)  # drop link URLs, keep text
    text = text.replace("`", "")
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def main() -> None:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    lines = MD.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue

        img = _IMG.match(line)
        if img:
            path = (HERE / img.group(2)).resolve()
            if path.exists():
                doc.add_picture(str(path), width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        elif line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").strip()) <= set("-: "):
            # Markdown table: header row, separator, then body rows.
            rows = [line]
            j = i + 2
            while j < len(lines) and lines[j].startswith("|"):
                rows.append(lines[j])
                j += 1
            cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
            table = doc.add_table(rows=len(cells), cols=len(cells[0]))
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(cells):
                for c, val in enumerate(row):
                    if c < len(table.rows[r].cells):
                        add_runs(table.rows[r].cells[c].paragraphs[0], val)
            i = j
            continue
        elif line.startswith("```"):
            j = i + 1
            buf = []
            while j < len(lines) and not lines[j].startswith("```"):
                buf.append(lines[j])
                j += 1
            p = doc.add_paragraph("\n".join(buf))
            p.style = doc.styles["Normal"]
            for run in p.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            i = j + 1
            continue
        elif line.startswith("> "):
            p = doc.add_paragraph()
            add_runs(p, line[2:])
            p.paragraph_format.left_indent = Inches(0.4)
        else:
            add_runs(doc.add_paragraph(), line)
        i += 1

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
