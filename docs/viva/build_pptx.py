"""Build the 12-slide viva deck AgentShield.pptx, using the real figures.

    pip install python-pptx
    python docs/viva/build_pptx.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
FIG = HERE.parent / "report" / "figures"
OUT = HERE / "AgentShield.pptx"

DARK = RGBColor(0x0E, 0x13, 0x20)
FG = RGBColor(0xE7, 0xEC, 0xF4)
ACCENT = RGBColor(0x5B, 0x8C, 0xFF)
MUTED = RGBColor(0x9A, 0xA6, 0xB8)


def _bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = DARK


def _text(slide, left, top, width, height, text, size, color=FG, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for k, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
    return box


def _bullets(slide, items, left=0.8, top=1.7, width=8.4, size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5))
    tf = box.text_frame
    tf.word_wrap = True
    for k, item in enumerate(items):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = "•  " + item
        run.font.size = Pt(size)
        run.font.color.rgb = FG
        p.space_after = Pt(10)


def title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    _text(slide, 0.8, 2.2, 8.4, 1.5, title, 40, ACCENT, bold=True)
    _text(slide, 0.8, 3.8, 8.4, 1.0, subtitle, 20, MUTED)
    return slide


def content_slide(prs, heading, items):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    _text(slide, 0.8, 0.6, 8.4, 0.9, heading, 30, ACCENT, bold=True)
    _bullets(slide, items)
    return slide


def figure_slide(prs, heading, fig, caption=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    _text(slide, 0.8, 0.6, 8.4, 0.9, heading, 30, ACCENT, bold=True)
    if (FIG / fig).exists():
        slide.shapes.add_picture(str(FIG / fig), Inches(1.6), Inches(1.7), width=Inches(6.8))
    else:
        _text(slide, 0.8, 3.0, 8.4, 1.0, f"[figure: {fig} — run eval/make_figures.py]", 16, MUTED)
    if caption:
        _text(slide, 0.8, 6.7, 8.4, 0.6, caption, 14, MUTED, align=PP_ALIGN.CENTER)
    return slide


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    title_slide(prs, "AgentShield", "Provenance-Aware Defense Against Multilingual\nPrompt Injection in LLM Agents  ·  [NAME], [GUIDE]")
    content_slide(prs, "The problem", [
        "LLM agents read web/email/files, then act: send mail, move money, call APIs.",
        "Indirect prompt injection: attacker hides instructions in that content.",
        "A hijack becomes a wrong, irreversible action — e.g. data exfiltration.",
    ])
    content_slide(prs, "Motivation", [
        "Agent-injection benchmarks are English-only.",
        "Low-resource & code-mixed inputs bypass safety training.",
        "Most defenses aren't tested against an adaptive attacker.",
    ])
    content_slide(prs, "Limits of existing defenses", [
        "Spotlighting/delimiters: rely on the model obeying the marking.",
        "Classifiers: English-centric and evadable.",
        "CaMeL/FIDES: strong guarantees but utility cost and complexity.",
        "None measure multilingual or code-mixed injection.",
    ])
    content_slide(prs, "AgentShield: the idea", [
        "Layered + provenance-aware: don't rely on detection alone.",
        "Multilingual detection (rules + xlm-roberta + LLM intent check).",
        "Deterministic backstop: canary tokens + taint tracking.",
        "Fused by logistic regression; enforced by a policy firewall.",
    ])
    content_slide(prs, "Architecture", [
        "One hook intercepts every tool call.",
        "Content: preprocess → detectors → fuse → sanitize/allow/block.",
        "Action: canary scan → taint policy → permissions/approval → allow/block.",
        "API + WebSocket + SDK + live dashboard + Docker.",
    ])
    content_slide(prs, "Detection pipeline", [
        "Preprocess: hidden HTML, base64/hex/URL/ROT13, NFKC, invisible chars, homoglyphs, OCR.",
        "L1 rules: multilingual triggers + structural signals (~0.1 ms).",
        "L2: fine-tuned xlm-roberta-base (recall layer).",
        "L3: LLM intent check (task vs. content).",
    ])
    content_slide(prs, "Canary & taint tracking (core contribution)", [
        "Canary: random token in prompt + sensitive files; any leak is blocked and traced.",
        "Taint: label data USER / PRIVATE / EXTERNAL.",
        "Policy: PRIVATE → EXTERNAL destination is blocked (or needs approval).",
        "Catches exfiltration that content detection misses.",
    ])
    content_slide(prs, "Adaptive attacker", [
        "Mutates seed attacks: encode, homoglyph, invisible, split, paraphrase, translate.",
        "Feedback by threat model: black / gray / white box.",
        "Reports attack success vs. query budget; feeds adversarial retraining.",
        "Finding: the preprocessor reverses obfuscations that evade a keyword filter.",
    ])
    figure_slide(prs, "Results & ablation", "comparison.png",
                 "Real numbers on a public English test set (L1 + fusion). See the report for details.")
    content_slide(prs, "Live demo", [
        "Unprotected agent leaks a fake secret from injected content.",
        "AgentShield on: the same attack is blocked, traced to the private file.",
        "Incident replay, multilingual catch, canary catch, dashboard on a phone.",
    ])
    content_slide(prs, "Future work & conclusion", [
        "Train L2 on GPU + add the human multilingual set → full headline numbers.",
        "Adversarial training; activation task-drift; strict dual-LLM upper bound.",
        "Takeaway: multilingual detection + provenance catches what detection misses.",
    ])

    prs.save(OUT)
    print(f"wrote {OUT} ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
